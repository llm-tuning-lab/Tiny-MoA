"""
Office Agent (Reference/Gemini-Claw 기반)
==========================================
PPT, Word, Excel 파일 생성을 위한 전문 모듈
한국어 폰트 지원 내장
"""

import os
from typing import List, Dict, Union, Optional
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment


class OfficeAgent:
    """
    Enhanced Office Automation Agent.
    Provides professional styling, structured document generation, and Korean font support.
    """
    def __init__(self, workspace_root: str = "."):
        self.workspace_root = os.path.abspath(workspace_root)
        self.default_font = "Malgun Gothic"  # 한국어 기본 폰트

    def _get_path(self, path: str) -> str:
        """경로를 절대 경로로 변환"""
        if os.path.isabs(path):
            return path
        return os.path.join(self.workspace_root, path)

    def _ensure_dir(self, path: str) -> None:
        """디렉토리가 없으면 생성"""
        dir_path = os.path.dirname(path)
        if dir_path and not os.path.exists(dir_path):
            os.makedirs(dir_path, exist_ok=True)

    def _apply_korean_font(self, run) -> None:
        """Helper to apply Korean font to a Word run."""
        run.font.name = self.default_font
        run._element.rPr.rFonts.set(qn('w:eastAsia'), self.default_font)

    # =========================================================================
    # Excel 생성
    # =========================================================================
    def create_excel(self, data: List[Dict], output_path: str, sheet_name: str = "Data") -> Dict:
        """
        Excel 파일 생성 (스타일링 포함)
        
        Args:
            data: 데이터 리스트 [{"Column1": "Value1", ...}, ...]
            output_path: 출력 파일 경로
            sheet_name: 시트 이름
        """
        try:
            full_path = self._get_path(output_path)
            self._ensure_dir(full_path)
            
            if not data:
                # 빈 데이터면 기본 데이터 생성
                data = [{"Info": "No data provided"}]
            
            df = pd.DataFrame(data)
            
            with pd.ExcelWriter(full_path, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                self._style_excel_sheet(writer.sheets[sheet_name], df)
            
            return {
                "success": True,
                "path": full_path,
                "message": f"Successfully created Excel with {len(data)} rows",
                "row_count": len(data)
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    def _style_excel_sheet(self, worksheet, dataframe, include_index=False):
        """Auto-adjust column widths and style headers with Korean font support."""
        header_font = Font(name=self.default_font, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        body_font = Font(name=self.default_font)
        
        for idx, col in enumerate(dataframe.columns):
            series = dataframe[col]
            try:
                max_len = max(
                    series.astype(str).map(len).max() if not series.empty else 0,
                    len(str(col))
                ) + 2
                worksheet.column_dimensions[get_column_letter(idx + 1)].width = min(max_len, 50)
            except:
                pass
            
            cell = worksheet.cell(row=1, column=idx + 1)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font

    # =========================================================================
    # Word 생성
    # =========================================================================
    def create_word_report(self, title: str, sections: List[Dict[str, str]], output_path: str) -> Dict:
        """
        Word 보고서 생성 (한국어 폰트 지원)
        
        Args:
            title: 문서 제목
            sections: 섹션 리스트 [{"heading": "...", "content": "..."}]
            output_path: 출력 파일 경로
        """
        try:
            full_path = self._get_path(output_path)
            self._ensure_dir(full_path)
            
            doc = Document()
            
            # 제목 스타일링
            title_heading = doc.add_heading(title, 0)
            title_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in title_heading.runs:
                self._apply_korean_font(run)
            
            # 섹션 추가
            for section in sections:
                heading_text = section.get("heading", "")
                content_text = section.get("content", "")
                
                if heading_text:
                    h = doc.add_heading(heading_text, level=1)
                    for run in h.runs:
                        self._apply_korean_font(run)
                
                if content_text:
                    for paragraph in content_text.split('\n'):
                        paragraph = paragraph.strip()
                        if not paragraph:
                            continue
                            
                        # Bullet Point 처리 (- 또는 * 로 시작하는 경우)
                        if paragraph.startswith("- ") or paragraph.startswith("* "):
                            p = doc.add_paragraph(paragraph[2:].strip(), style='List Bullet')
                        else:
                            p = doc.add_paragraph(paragraph)
                            
                        # 폰트 적용 (스타일 적용 후에도 개별 런에 폰트 적용 필요)
                        for run in p.runs:
                            self._apply_korean_font(run)
            
            doc.save(full_path)
            return {
                "success": True,
                "path": full_path,
                "message": f"Successfully created Word with {len(sections)} sections",
                "section_count": len(sections)
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    # =========================================================================
    # PowerPoint 생성
    # =========================================================================
    def create_presentation(self, title: str, subtitle: str, slides: List[Dict], output_path: str) -> Dict:
        """
        PowerPoint 프레젠테이션 생성 (한국어 폰트 지원)
        
        Args:
            title: 발표 제목
            subtitle: 부제목
            slides: 슬라이드 데이터 [{"title": "...", "content": ["...", "..."]}]
            output_path: 출력 파일 경로
        """
        try:
            full_path = self._get_path(output_path)
            self._ensure_dir(full_path)
            
            prs = Presentation()
            
            # 1. Title Slide (Layout 0)
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            slide.placeholders[1].text = subtitle
            
            # Apply fonts to title slide
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = self.default_font
                    for run in paragraph.runs:
                        run.font.name = self.default_font

            # 2. Content Slides (Layout 1)
            bullet_layout = prs.slide_layouts[1]
            
            for slide_info in slides:
                s_title = slide_info.get("title", "Untitled")
                s_content = slide_info.get("content", [])
                
                slide = prs.slides.add_slide(bullet_layout)
                slide.shapes.title.text = s_title
                
                tf = slide.placeholders[1].text_frame
                
                if isinstance(s_content, list):
                    for i, point in enumerate(s_content):
                        if i == 0:
                            tf.text = str(point)
                        else:
                            p = tf.add_paragraph()
                            p.text = str(point)
                elif isinstance(s_content, str):
                    tf.text = s_content

                # Apply fonts to all text in slide
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        paragraph.font.name = self.default_font
                        for run in paragraph.runs:
                            run.font.name = self.default_font
            
            prs.save(full_path)
            return {
                "success": True,
                "path": full_path,
                "message": f"Successfully created presentation with {len(slides)} slides",
                "slide_count": len(slides)
            }
        except Exception as e:
            return {"success": False, "error": str(e)}
